"""Caracterização das fronteiras antes das extrações; sem serviços externos."""
import asyncio
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, Mock, patch

from fastapi import HTTPException
from playwright.async_api import TimeoutError as PlaywrightTimeoutError
from pydantic import ValidationError

from app import captcha
from app.api import _safe_call, health
from app.consultar import _limpar_cnpj
from app.models import ConsultaRequest
from scripts import batch_optantes_xlsx as batch


class InputContractTests(unittest.TestCase):
    def test_mascara_e_digitos_unicode_mantem_aceitacao_atual(self):
        for value in ('43.996.591/0001-01', '43996591000101', '４３９９６５９１０００１０１'):
            with self.subTest(value=value):
                self.assertEqual(ConsultaRequest(cnpj=value).cnpj, value)
                self.assertEqual(_limpar_cnpj(value), value.replace('.', '').replace('/', '').replace('-', ''))

    def test_entrada_invalida_preserva_mensagens_por_fronteira(self):
        for value in ('', '123', ' 43996591000101', 'ABCDEFGHIJKLMN'):
            with self.subTest(value=value):
                with self.assertRaises(ValidationError):
                    ConsultaRequest(cnpj=value)
                with self.assertRaisesRegex(ValueError, 'CNPJ invalido:'):
                    _limpar_cnpj(value)


class ErrorContractTests(unittest.IsolatedAsyncioTestCase):
    async def test_mapeamento_http_e_mensagens_existentes(self):
        for error, status in ((ValueError('entrada'), 400), (RuntimeError('origem'), 502),
                              (TimeoutError('prazo'), 504), (PlaywrightTimeoutError('browser'), 504),
                              (Exception('inesperado'), 500)):
            with self.subTest(status=status, kind=type(error)):
                with self.assertRaises(HTTPException) as raised:
                    await _safe_call(AsyncMock(side_effect=error))
                self.assertEqual(raised.exception.status_code, status)
                self.assertIn(str(error), raised.exception.detail)

    async def test_health_independe_de_browser_e_captcha(self):
        with patch('app.browser.create_browser', side_effect=AssertionError('browser')), \
             patch.object(captcha.requests, 'post', side_effect=AssertionError('solver')):
            self.assertEqual(await health(), {'status': 'ok'})


class CaptchaContractTests(unittest.TestCase):
    def response(self, payload, status=200, text='json'):
        return Mock(status_code=status, text=text, json=Mock(return_value=payload))

    def test_nonecap_rejeita_chave_ausente_sem_http(self):
        with patch.object(captcha, 'NONECAP_API_KEY', ''), patch.object(captcha.requests, 'post') as request:
            with self.assertRaisesRegex(RuntimeError, 'nao configurada'):
                captcha.resolver_hcaptcha_nonecap('https://local.test', 'site')
            request.assert_not_called()

    def test_nonecap_deadline_e_respostas_invalidas(self):
        with patch.object(captcha, 'NONECAP_API_KEY', 'teste'), patch.object(captcha.time, 'monotonic', side_effect=[0, 2]):
            with self.assertRaises(TimeoutError):
                captcha.resolver_hcaptcha_nonecap('https://local.test', 'site', timeout_s=1)
        for payload in ([], {'status': 'solved'}, {'status': 'pending', 'id': 'invalido'}):
            with self.subTest(payload=payload), patch.object(captcha, 'NONECAP_API_KEY', 'teste'), \
                 patch.object(captcha.requests, 'post', return_value=self.response(payload)):
                with self.assertRaises(RuntimeError):
                    captcha.resolver_hcaptcha_nonecap('https://local.test', 'site')

    def test_nonecap_timeout_transporte_e_json_invalido(self):
        for effect, expected in ((captcha.requests.Timeout(), TimeoutError), (ValueError('json'), RuntimeError)):
            response = self.response({})
            response.json.side_effect = effect
            with patch.object(captcha, 'NONECAP_API_KEY', 'teste'), patch.object(captcha.requests, 'post', return_value=response):
                with self.assertRaises(expected):
                    captcha.resolver_hcaptcha_nonecap('https://local.test', 'site')

    def test_nopecha_submit_poll_e_estados_de_espera(self):
        pending = self.response({'error': 14}, status=409)
        solved = self.response({'data': 'T' * 101})
        with patch.object(captcha, 'NOPECHA_API_KEY', 'teste'), patch.object(captcha.time, 'sleep'), \
             patch.object(captcha.requests, 'post', return_value=self.response({'data': 'job'})), \
             patch.object(captcha.requests, 'get', side_effect=[self.response({}, text=''), pending, solved]):
            self.assertEqual(captcha.resolver_hcaptcha('https://local.test', 'site'), 'T' * 101)

    def test_nopecha_erros_e_timeout_preservam_tipos(self):
        with patch.object(captcha, 'NOPECHA_API_KEY', ''):
            with self.assertRaisesRegex(RuntimeError, 'nao configurada'):
                captcha.resolver_hcaptcha('https://local.test', 'site')
        with patch.object(captcha, 'NOPECHA_API_KEY', 'teste'), patch.object(captcha.requests, 'post', return_value=self.response({}, text='')):
            with self.assertRaisesRegex(RuntimeError, 'resposta vazia'):
                captcha.resolver_hcaptcha('https://local.test', 'site')
        with self.assertRaisesRegex(RuntimeError, 'NoPeCHA timeout'):
            captcha._poll_token('job', timeout_s=0)


class ConfigContractTests(unittest.TestCase):
    def test_env_da_raiz_e_precedencia_do_ambiente(self):
        # Executa a configuração numa raiz temporária: nunca abre o .env real.
        source = Path(__file__).resolve().parents[1] / 'app' / 'config.py'
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'app').mkdir()
            target = root / 'app' / 'config.py'
            target.write_bytes(source.read_bytes())
            (root / '.env').write_text('NONECAP_API_KEY=arquivo\nNOPECHA_API_KEY=arquivo\n', encoding='utf8')
            with patch.dict(os.environ, {'NONECAP_API_KEY': 'ambiente', 'NOPECHA_API_KEY': '', 'PYTHON_DOTENV_DISABLED': '0'}):
                os.environ.pop('NOPECHA_API_KEY')
                spec = importlib.util.spec_from_file_location('config_isolada', target)
                module = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(module)
                self.assertEqual(module.NONECAP_API_KEY, 'ambiente')
                self.assertEqual(module.NOPECHA_API_KEY, 'arquivo')


class BatchContractTests(unittest.IsolatedAsyncioTestCase):
    async def test_retomada_aceita_cache_antigo_e_repete_apenas_falhas(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            out = root / 'output'
            out.mkdir()
            cache = out / 'optantes_batch_20000101_0000.json'
            cache.write_text(json.dumps({'results': {'00000000000001': {'data': {}, 'erro': None},
                                                     '00000000000002': {'data': None, 'erro': 'falhou'}}}), encoding='utf8')
            consult = AsyncMock(side_effect=RuntimeError('indisponivel'))
            with patch.object(batch, 'ROOT', root), patch.object(batch, 'CNPJS', ['1', '2']), \
                 patch('app.consultar.consultar_apenas_optantes', consult), patch.object(batch, '_write_xlsx'):
                await batch.main()
            consult.assert_awaited_once_with('00000000000002')
            saved = json.loads(cache.read_text(encoding='utf8'))['results']
            self.assertEqual(saved['00000000000001'], {'data': {}, 'erro': None})
            self.assertEqual(saved['00000000000002']['erro'], 'indisponivel')

    async def test_exportacao_preserva_abas_e_celulas(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'resultado.xlsx'
            batch._write_xlsx([batch._row_from_result('00000000000001', None, 'erro')], target)
            workbook = load_workbook(target)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook.sheetnames, ['Todas', 'Optantes Simples', 'MEI', 'Eventos Futuros', 'Erros'])
            self.assertEqual(workbook['Erros']['A2'].value, '00000000000001')
            self.assertEqual(workbook['Erros']['L2'].value, 'erro')
